"""Tests for the Office loader (Task 2.6) -- ``.docx`` and ``.pptx``.

``office_loader`` extracts:

* ``.docx`` via ``python-docx``: paragraph texts concatenated into ``text``.
* ``.pptx`` via ``python-pptx``: ONE logical unit per slide -- all shape text plus
  the slide's speaker notes -- joined with a per-slide separator; ``n_slides`` in
  ``native_metadata``.

``source_type`` comes from :func:`bios667_rag.metadata.classify_source_type`
(``.pptx`` -> ``slides``; ``.docx`` -> role-based).

Fixtures are built on the fly with ``python-docx``/``python-pptx`` so the tests
carry no binary blobs and run anywhere those libraries are installed.
"""

from docx import Document
import pytest
from pptx import Presentation

from bios667_rag.loaders import load
from bios667_rag.loaders.office import office_loader
from bios667_rag.metadata import SOURCE_TYPES


def _make_docx(path, paragraphs):
    """Create a ``.docx`` whose body is ``paragraphs`` (a list of strings)."""
    d = Document()
    for para in paragraphs:
        d.add_paragraph(para)
    d.save(str(path))


def _make_pptx(path, slides):
    """Create a ``.pptx`` from ``slides``: list of ``(title, notes_or_None)``."""
    p = Presentation()
    for title, notes in slides:
        s = p.slides.add_slide(p.slide_layouts[5])  # "Title Only" layout
        s.shapes.title.text = title
        if notes is not None:
            s.notes_slide.notes_text_frame.text = notes
    p.save(str(path))


# --- .docx ---------------------------------------------------------------


def test_docx_concatenates_paragraphs(tmp_path):
    f = tmp_path / "BIOS_667_2025_syllabus.docx"
    _make_docx(f, ["Longitudinal data analysis syllabus", "Mixed models week 5"])

    rd = office_loader(str(f))

    assert rd is not None
    assert "Longitudinal data analysis syllabus" in rd.text
    assert "Mixed models week 5" in rd.text


def test_docx_drops_empty_paragraph_runs(tmp_path):
    f = tmp_path / "BIOS_667_2025_syllabus.docx"
    # Word commonly inserts empty paragraphs as vertical spacing; interleave them.
    _make_docx(f, ["First line", "", "   ", "Second line", ""])

    rd = office_loader(str(f))

    assert rd is not None
    assert "First line" in rd.text
    assert "Second line" in rd.text
    # No blank-line runs: no consecutive newlines anywhere in the extracted text.
    assert "\n\n" not in rd.text
    assert rd.text == "First line\nSecond line"


def test_docx_source_type_is_role_based(tmp_path):
    f = tmp_path / "BIOS_667_2025_syllabus.docx"
    _make_docx(f, ["Syllabus body"])

    rd = office_loader(str(f))

    # "syllabus" in filename -> classified as syllabus.
    assert rd.source_type == "syllabus"
    assert rd.source_type in SOURCE_TYPES


# --- .pptx ---------------------------------------------------------------


def test_pptx_captures_titles_and_notes(tmp_path):
    f = tmp_path / "lecture12_marginal.pptx"
    _make_pptx(
        f,
        [
            ("GEE overview", "remember MCAR"),
            ("Working correlation", None),
        ],
    )

    rd = office_loader(str(f))

    assert rd is not None
    # Slide title text present.
    assert "GEE overview" in rd.text
    assert "Working correlation" in rd.text
    # Speaker-notes text present.
    assert "MCAR" in rd.text
    assert rd.source_type == "slides"
    assert rd.native_metadata["n_slides"] == 2


# --- dispatch integration ------------------------------------------------


def test_dispatch_loads_docx(tmp_path):
    f = tmp_path / "notes.docx"
    _make_docx(f, ["Random intercept model"])

    # Proves registration + BINARY_EXTS skip: a binary zip-based file would be
    # rejected by dispatch's UTF-8 pre-read if .docx were not in BINARY_EXTS.
    rd = load(str(f))

    assert rd is not None
    assert "Random intercept model" in rd.text


def test_dispatch_loads_pptx(tmp_path):
    f = tmp_path / "deck.pptx"
    _make_pptx(f, [("Linear mixed effects", None)])

    rd = load(str(f))

    assert rd is not None
    assert "Linear mixed effects" in rd.text
    assert rd.source_type == "slides"
