"""Loader for Microsoft Office documents: ``.docx`` and ``.pptx``.

Both formats are zip-based binaries (not UTF-8 text), so their extensions are listed
in :data:`bios667_rag.loaders.dispatch.BINARY_EXTS` to skip dispatch's UTF-8 decode
pre-read; the loader itself owns parse/validity (returning ``None`` to reject a file).

* ``.docx`` (Word) via ``python-docx``: every non-empty body paragraph's text is
  concatenated into ``text`` (one paragraph per line). Empty/whitespace-only
  paragraphs are dropped so blank-line runs don't degrade embedding quality.
  Tables are intentionally not flattened -- paragraph text is sufficient for the
  course corpus (syllabi, handouts, prose).

* ``.pptx`` (PowerPoint) via ``python-pptx``: ONE logical unit per slide. For each
  slide we gather text from every shape that has a text frame, PLUS the slide's
  speaker notes (when present), and join slides with a ``--- Slide N ---`` separator
  so downstream chunking can see slide boundaries. ``native_metadata["n_slides"]``
  records the slide count.

``source_type`` is resolved by :func:`bios667_rag.metadata.classify_source_type`
(``.pptx`` -> ``slides``; ``.docx`` -> role-based, e.g. ``syllabus`` by filename).

A corrupt/unparseable file returns ``None``, consistent with the loader contract.
"""

from pathlib import Path

from docx import Document
from pptx import Presentation

from ..metadata import classify_source_type
from .base import RawDoc
from .dispatch import register


def _load_docx(path: str) -> RawDoc | None:
    """Extract Word paragraph text into a :class:`RawDoc`, or ``None`` if corrupt."""
    try:
        doc = Document(path)
    except Exception:
        return None

    # Skip empty/whitespace-only paragraphs so blank-line runs don't degrade
    # embedding quality (Word inserts many empty paragraphs as vertical spacing).
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())

    return RawDoc(
        text=text,
        source_path=path,
        source_type=classify_source_type(path),
        native_metadata={},
    )


def _slide_text(slide) -> str:
    """Gather all shape text plus speaker notes for one PowerPoint slide."""
    parts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            txt = shape.text_frame.text
            if txt:
                parts.append(txt)
    # Speaker notes live on a separate notes slide that may not exist.
    if slide.has_notes_slide:
        notes = slide.notes_slide.notes_text_frame.text
        if notes:
            parts.append(notes)
    return "\n".join(parts)


def _load_pptx(path: str) -> RawDoc | None:
    """Extract per-slide text + notes into a :class:`RawDoc`, or ``None`` if corrupt."""
    try:
        prs = Presentation(path)
    except Exception:
        return None

    slides = list(prs.slides)
    blocks = [
        f"--- Slide {idx} ---\n{_slide_text(slide)}"
        for idx, slide in enumerate(slides, start=1)
    ]
    text = "\n\n".join(blocks)

    return RawDoc(
        text=text,
        source_path=path,
        source_type=classify_source_type(path),
        native_metadata={"n_slides": len(slides)},
    )


def office_loader(path: str) -> RawDoc | None:
    """Load a ``.docx`` or ``.pptx`` file into a :class:`RawDoc`, or ``None``.

    Dispatch routes by extension, but this loader also branches on the suffix so it
    can be called directly. An unexpected extension (should not happen via dispatch)
    returns ``None``.
    """
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".pptx":
        return _load_pptx(path)
    return None


register(".docx", ".pptx", loader=office_loader)
